import asyncio
import io
import json
import pathlib
import re
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from google.adk.runners import InMemoryRunner
from google.genai import types as genai_types

from agent import a2ui, agent as dash, analysis, charts, llm, page, pdf, render, sheets, validate
from agent.schema import Insight, Narrative

SAMPLE = pathlib.Path(__file__).resolve().parent.parent / "samples" / "sample_workbook.xlsx"
BOOK = sheets.Upload("sample_workbook.xlsx", SAMPLE.read_bytes(), sheets.XLSX)
CSV_TEXT = ("Quarterly results\n\nRegion,Q1 2024,Q2 2024,Q3 2024,Q4 2024\nNorth,\"£1,200\",\"£1,350\",\"£1,410\",\"£1,600\"\n"
            "South,900,950,(100),1020\nEast,400,420,480,510\nTotal,2500,2720,1790,3130\n")


@pytest.fixture(scope="module")
def tabs():
    return sheets.read(BOOK)


def _tab(tabs, name):
    return next(t for t in tabs if t.name == name)


# --- reading ----------------------------------------------------------------------------

def test_reads_every_tab_and_tidies(tabs):
    assert [t.name for t in tabs] == ["Monthly Sales", "Spend by Dept", "Customers"]
    spend = _tab(tabs, "Spend by Dept")
    assert spend.wide and len(spend.notes) == 3  # title row skipped, Total row dropped, wide -> long
    assert "Total" not in set(spend.df["Department"])
    p = sheets.profile(spend)
    assert (p.time, p.grain, p.measures, p.dimensions) == ("Period", "month", ["Value"], ["Department"])
    sales = sheets.profile(_tab(tabs, "Monthly Sales"))
    assert sales.time == "Date" and sales.measures[0] == "Revenue" and set(sales.dimensions) == {"Region", "Product"}
    cust = sheets.profile(_tab(tabs, "Customers"))
    assert "Customer ID" not in cust.measures + cust.dimensions and cust.time is None


def test_reads_csv_with_title_currency_brackets_and_quarters():
    tab, = sheets.read(sheets.Upload("results.csv", CSV_TEXT.encode(), sheets.CSV))
    assert tab.name == "results" and tab.wide
    p = sheets.profile(tab)
    assert p.grain == "quarter" and p.span == "Q1 2024 to Q4 2024"
    north = tab.df[tab.df["Region"] == "North"]["Value"].tolist()
    assert north == [1200, 1350, 1410, 1600]
    assert -100 in tab.df["Value"].tolist()


@pytest.mark.parametrize("raw,expected", [("Jan 2024", "2024-01-01"), ("Q3 2024", "2024-07-01"), ("2024 Q2", "2024-04-01"),
                                          ("FY2023", "2023-01-01"), ("2024-11", "2024-11-01"), ("Sept 24", "2024-09-01"), ("hello", None)])
def test_period_of(raw, expected):
    got = sheets.period_of(raw)
    assert (got.strftime("%Y-%m-%d") if got is not None else None) == expected


@pytest.mark.parametrize("raw,expected", [("£1,200", 1200), ("(300)", -300), ("12%", 12), ("1.5k", 1500), ("USD40", 40), ("18°C", 18), ("-3.5°F", -3.5), ("abc", None), (True, None)])
def test_to_number(raw, expected):
    assert sheets.to_number(raw) == expected


def test_uploads_skip_our_own_datapart_and_non_spreadsheets():
    wrapped = a2ui.to_genai_part({"version": "v0.9"})
    pdf_part = genai_types.Part(inline_data=genai_types.Blob(mime_type="application/pdf", data=b"%PDF", display_name="a.pdf"))
    csv_part = genai_types.Part(inline_data=genai_types.Blob(mime_type="application/octet-stream", data=b"a,b\n1,2", display_name="x.csv"))
    ups = sheets.uploads_in(genai_types.Content(role="user", parts=[wrapped, pdf_part, csv_part]))
    assert [(u.name, u.mime) for u in ups] == [("x.csv", sheets.CSV)]


# --- analysis ---------------------------------------------------------------------------

@pytest.mark.parametrize("tab,style", [(t, s) for t in ("Monthly Sales", "Spend by Dept", "Customers") for s in analysis.STYLES])
def test_every_style_on_every_tab_renders_and_validates(tabs, tab, style):
    a = analysis.analyze(_tab(tabs, tab), style)
    assert a.style in analysis.STYLES and a.style != "auto"
    assert a.facts and len(a.kpis) == 4 and a.chart["type"] == a.style and a.table["rows"]
    messages = render.dashboard_card(a, llm.fallback(a), BOOK.name)
    validate.check(messages)


def test_style_that_does_not_fit_falls_back_with_a_note(tabs):
    a = analysis.analyze(_tab(tabs, "Customers"), "trend")
    assert a.style == "explore" and any("needs a date" in n for n in a.notes)


def test_sales_facts_find_partial_month_anomalies_and_drivers(tabs):
    a = analysis.analyze(_tab(tabs, "Monthly Sales"), "trend")
    text = a.fact_lines()
    assert "Dec 2025 is only partly in the data" in text
    assert a.chart["labels"][-1] == "Nov 2025"  # the part-month is left out of the trend
    assert "41.2K · East · Alpha · 12 Mar 2025" in text and "23.9K · North · Beta · 05 Aug 2024" in text  # planted
    assert "seasonal pattern: November" in text
    assert "Revenue moves strongly with Units" in text
    assert len(a.chart["forecast"]) == 3 and all(f["lo"] < f["value"] < f["hi"] for f in a.chart["forecast"])


def test_bridge_steps_add_up_and_heatmap_flags_the_spike(tabs):
    spend = _tab(tabs, "Spend by Dept")
    b = analysis.analyze(spend, "bridge").chart
    assert b["start"]["value"] + sum(s["value"] for s in b["steps"]) == pytest.approx(b["end"]["value"])
    h = analysis.analyze(spend, "heatmap")
    assert any("Marketing in Sep 2025" in f.text for f in h.facts)


def test_pareto_cut_marks_80_percent(tabs):
    c = analysis.analyze(_tab(tabs, "Monthly Sales"), "pareto", dim="Region").chart
    assert c["items"][c["cut"]]["cum"] >= 80 and (c["cut"] == 0 or c["items"][c["cut"] - 1]["cum"] < 80)


def test_num_format():
    assert [analysis.num(v) for v in (1234567, 15320, 912.456, 12, 3.14159, -85_800)] == ["1.23M", "15.3K", "912", "12", "3.14", "-85.8K"]


def test_nice_ticks():
    assert charts.nice_ticks(0, 97) == [0, 25, 50, 75, 100]


# --- narrative --------------------------------------------------------------------------

def test_numbers_parse_suffixes_and_ignore_ids():
    assert llm.numbers("went from 85.8K to 1.2M (+14.0%), 2.4x, Q3 and F12") == [85800, 1_200_000, 14.0, 2.4]


def test_grounded_replaces_invented_numbers(tabs):
    a = analysis.analyze(_tab(tabs, "Monthly Sales"), "trend")
    real = next(f for f in a.facts if f.kind == "trend")
    n = Narrative(headline="Revenue will hit 999.9M next year", summary=real.text,
                  insights=[Insight(kind="trend", title="Up a lot", text=real.text, evidence=[real.id]),
                            Insight(kind="driver", title="Made up", text="North grew 777.7% last week.", evidence=[real.id, "F999"]),
                            Insight(kind="risk", title="No evidence", text="Costs rose 555.5K.", evidence=[])],
                  actions=["Cut costs by 31.4159%", "Check the flagged rows"], questions=["Which Region grew fastest?"])
    g = llm.grounded(n, a)
    assert "999.9" not in g.headline and g.summary == real.text
    assert g.insights[0].text == real.text
    assert g.insights[1].text == real.text and g.insights[1].evidence == [real.id]  # replaced by its evidence
    assert all("555.5" not in i.text for i in g.insights)
    assert g.actions == ["Check the flagged rows"]


def test_prompt_carries_facts_columns_and_focus(tabs):
    a = analysis.analyze(_tab(tabs, "Monthly Sales"), "heatmap", dim="Product")
    text = llm.prompt(a, BOOK.name, "Why did West fall?")
    assert "F1 [summary]" in text and "Product (category)" in text and 'The user asked: "Why did West fall?"' in text


# --- the agent end to end ---------------------------------------------------------------

class FakeModels:
    """Answers like Gemini: a Narrative that cites the first facts in the prompt."""

    def __init__(self, fail=False):
        self.calls, self.fail = [], fail

    async def generate_content(self, model, contents, config):
        self.calls.append((contents, config))
        if self.fail:
            raise RuntimeError("quota")
        facts = re.findall(r"^(F\d+) \[(\w+)\] (.*)$", contents, re.M)
        picked = [f for f in facts if f[1] not in ("summary", "quality")][:3]
        n = Narrative(headline=picked[0][2][:100], summary=" ".join(f[2] for f in picked[:2]),
                      insights=[Insight(kind="trend", title="What happened", text=f[2], evidence=[f[0]]) for f in picked],
                      actions=["Check the flagged rows."], questions=["Which Region grew fastest?", "Show it as a heatmap", "What about Units?"])
        return SimpleNamespace(parsed=n, text=n.model_dump_json())


@pytest.fixture
def fake(monkeypatch):
    models = FakeModels()
    monkeypatch.setattr(dash, "_client", SimpleNamespace(aio=SimpleNamespace(models=models)))
    return models


def _upload(up: sheets.Upload, text: str | None = None):
    parts = [genai_types.Part(inline_data=genai_types.Blob(mime_type=up.mime, data=up.data, display_name=up.name))]
    if text:
        parts.append(genai_types.Part(text=text))
    return genai_types.Content(role="user", parts=parts)


def _text(text: str):
    return genai_types.Content(role="user", parts=[genai_types.Part(text=text)])


def _click(name: str, context: dict):
    msg = {"version": "v0.9", "action": {"name": name, "context": context, "sourceComponentId": "x", "surfaceId": "s", "timestamp": "t"}}
    blob = a2ui.TAG_START + json.dumps({"kind": "data", "data": msg, "metadata": {"mimeType": a2ui.A2UI_MIME_TYPE}}).encode() + a2ui.TAG_END
    return genai_types.Content(role="user", parts=[genai_types.Part(inline_data=genai_types.Blob(mime_type="text/plain", data=blob))])


def _messages(events) -> list[dict]:
    out = []
    for e in events:
        for p in e.content.parts:
            if p.inline_data and p.inline_data.data.startswith(a2ui.TAG_START):
                out.append(json.loads(p.inline_data.data[len(a2ui.TAG_START):-len(a2ui.TAG_END)])["data"])
    return out


def _components(messages) -> dict:
    return {c["id"]: c for m in messages if "updateComponents" in m for c in m["updateComponents"]["components"]}


def _converse(turns):
    async def run():
        runner = InMemoryRunner(agent=dash.root_agent, app_name="dash")
        session = await runner.session_service.create_session(app_name="dash", user_id="u")
        out = []
        for msg in turns:
            out.append([e async for e in runner.run_async(user_id="u", session_id=session.id, new_message=msg)])
        return out
    return asyncio.run(run())


def test_upload_tab_style_dashboard_ask_and_pdf(fake):
    turns = _converse([
        _upload(BOOK),
        _click("tab", {"sheet": "Spend by Dept", "prompt": "Use the Spend by Dept tab"}),
        _click("chart", {"sheet": "Spend by Dept", "style": "bridge", "prompt": "Show Spend by Dept as change bridge"}),
        _click("ask", {"sheet": "Spend by Dept", "style": "bridge", "measure": "Value", "dim": "Department", "prompt": "Why did Marketing jump?"}),
        _click("pdf", {"sheet": "Spend by Dept", "style": "bridge", "measure": "Value", "dim": "Department"}),
    ])
    tabs_c = _components(_messages(turns[0]))
    assert {"tab0", "tab1", "tab2"} <= set(tabs_c) and tabs_c["tab1"]["action"]["event"]["context"]["sheet"] == "Spend by Dept"
    styles_c = _components(_messages(turns[1]))
    assert "st-bridge" in styles_c and "st-scatter" not in styles_c  # one number column: no relationship chart
    dash_msgs = _messages(turns[2])
    validate.check(dash_msgs)
    comps = _components(dash_msgs)
    assert comps["root"]["component"] == "Canvas" and "What moved Value" in comps["frame"]["htmlContent"]
    assert len(fake.calls) == 2 and "Why did Marketing jump?" in fake.calls[1][0]
    assert _components(_messages(turns[3]))["root"]["component"] == "Canvas"
    files = [p for e in turns[4] for p in e.content.parts if p.inline_data and p.inline_data.mime_type == "application/pdf"]
    assert len(files) == 1 and files[0].inline_data.data.startswith(b"%PDF") and files[0].inline_data.display_name == "Spend-by-Dept-insights.pdf"
    assert len(fake.calls) == 2  # the PDF needs no model call


def test_typed_request_with_upload_goes_straight_to_the_chart(fake):
    first, = _converse([_upload(BOOK, "Show Revenue by Product in Monthly Sales as a heatmap")])
    html_ = _components(_messages(first))["frame"]["htmlContent"]
    assert "Revenue by Product and Date" in html_
    assert first[0].content.parts[0].text.startswith("**")


def test_follow_up_question_keeps_the_chart_and_passes_focus(fake):
    turns = _converse([_upload(BOOK, "Use Monthly Sales as small multiples"), _text("Why did West fall?")])
    assert "one panel per Region" in _components(_messages(turns[1]))["frame"]["htmlContent"]
    assert 'The user asked: "Why did West fall?"' in fake.calls[-1][0]


def test_model_failure_still_ships_a_facts_only_dashboard(monkeypatch):
    monkeypatch.setattr(dash, "_client", SimpleNamespace(aio=SimpleNamespace(models=FakeModels(fail=True))))
    first, = _converse([_upload(BOOK, "Customers as a scatter")])
    msgs = _messages(first)
    validate.check(msgs)
    assert "Annual Spend against Tenure (months)" in _components(msgs)["frame"]["htmlContent"]


def test_no_upload_and_wrong_file_type_get_the_welcome(fake):
    hello, = _converse([_text("hi")])
    assert "dash-welcome" in _messages(hello)[0]["createSurface"]["surfaceId"]
    wrong, = _converse([_upload(sheets.Upload("deck.pdf", b"%PDF-1.4", "application/pdf"))])
    assert "Excel (.xlsx) and CSV" in wrong[0].content.parts[0].text


def test_single_tab_csv_skips_the_tab_picker(fake):
    first, = _converse([_upload(sheets.Upload("results.csv", CSV_TEXT.encode(), sheets.CSV))])
    assert "dash-styles" in _messages(first)[0]["createSurface"]["surfaceId"]


def test_page_is_self_contained(tabs):
    a = analysis.analyze(_tab(tabs, "Monthly Sales"), "trend")
    html_ = page.page(a, llm.fallback(a), BOOK.name)
    assert "<script src" not in html_ and "http" not in html_.split("<script>")[0].replace("http-equiv", "")
    assert html_.count('data-act="chart"') >= 6 and "postMessage" in html_


def test_pdf_builds_for_every_style(tabs):
    for style in analysis.STYLES:
        a = analysis.analyze(_tab(tabs, "Monthly Sales"), style)
        assert pdf.build(a, llm.fallback(a), BOOK.name).startswith(b"%PDF")


# --- interactive explorer ---------------------------------------------------------------

def _explorer_data(html_: str) -> dict:
    m = re.search(r'<script type="application/json" id="ex-data">(.*?)</script>', html_, re.S)
    return json.loads(m.group(1).replace("<\\/", "</"))


def test_treemap_groups_by_the_coarse_category_and_colours_by_a_score(tabs):
    a = analysis.analyze(_tab(tabs, "Customers"), "explore")
    d, dims = a.chart["defaults"], a.chart["cube"]["dims"]
    assert dims[d["tmGroup"]]["name"] == "Segment" and dims[d["tmCell"]]["name"] == "Customer ID"
    assert a.chart["cube"]["measures"][int(d["tmColor"].split(":")[1])]["name"] == "NPS score"
    sales = analysis.analyze(_tab(tabs, "Monthly Sales"), "explore").chart
    assert sales["cube"]["dims"][sales["defaults"]["tmGroup"]]["name"] == "Product"  # 4 products group 5 regions


def test_auto_picks_the_explorer_when_there_are_categories(tabs):
    for name in ("Monthly Sales", "Spend by Dept", "Customers"):
        assert analysis.analyze(_tab(tabs, name)).style == "explore"


@pytest.mark.parametrize("tab", ["Monthly Sales", "Spend by Dept", "Customers"])
def test_explorer_cube_page_validates_and_fits(tabs, tab):
    a = analysis.analyze(_tab(tabs, tab), "explore")
    messages = render.dashboard_card(a, llm.fallback(a), BOOK.name)
    validate.check(messages)
    html_ = _components(messages)["frame"]["htmlContent"]
    assert len(json.dumps(messages).encode()) < 800_000
    data = _explorer_data(html_)
    cube, d = data["cube"], data["defaults"]
    width = len(cube["dims"]) + 2 + 2 * len(cube["measures"])
    assert cube["rows"] and all(len(r) == width for r in cube["rows"])
    for key in ("x", "y", "size"):  # every default metric is one the page offers
        mid = d[key]
        assert mid == "rows" or int(mid.split(":")[1]) < len(cube["measures"])
    for vid in ("ex-v-bubble", "ex-v-tree", "ex-v-heat", "ex-v-table", "ex-side", "ex-rank", "ex-card"):  # four views, side story, ranked list
        assert f'id="{vid}"' in html_
    dims = cube["dims"]
    assert d["tmCell"] < len(dims) and (d["tmGroup"] is None or (d["tmGroup"] < len(dims) and d["tmGroup"] != d["tmCell"]))
    assert d["tmSize"] == "rows" or not cube["measures"][int(d["tmSize"].split(":")[1])]["rate"]  # blocks are sized by an amount


def test_cube_totals_match_the_sheet(tabs):
    sheet = _tab(tabs, "Spend by Dept")
    a = analysis.analyze(sheet, "explore")
    cube = _explorer_data(page.page(a, llm.fallback(a), BOOK.name))["cube"]
    nd, i = len(cube["dims"]), next(k for k, m in enumerate(cube["measures"]) if m["name"] == a.measure)
    in_cube = sum(r[nd + 2 + 2 * i] for r in cube["rows"])
    df = sheet.df
    assert in_cube == pytest.approx(df[a.measure].sum(), rel=0.02)  # less only the partial period, if any


def test_filters_narrow_the_rows_and_are_described(tabs):
    sheet = _tab(tabs, "Monthly Sales")
    a = analysis.analyze(sheet, "explore", filters={"dims": {"Region": ["South"], "Nope": ["x"]}, "from": "2024-06-01", "to": "2024-06-30"})
    assert a.filters == {"dims": {"Region": ["South"]}, "from": "2024-06-01", "to": "2024-06-30"}
    assert a.meta()["filters"] == a.filters
    assert any("filtered to Region = South" in f.text for f in a.facts)
    with pytest.raises(ValueError, match="no rows match"):
        analysis.analyze(sheet, "explore", filters={"dims": {"Region": ["Atlantis"]}})


def test_explain_this_view_reanalyses_the_filtered_rows(fake):
    f = {"dims": {"Region": ["South"]}, "from": "2024-06-01", "to": "2024-06-30"}
    turns = _converse([
        _upload(BOOK, "Use Monthly Sales as interactive explorer"),
        _click("ask", {"sheet": "Monthly Sales", "style": "explore", "measure": "Revenue", "dim": "Region", "filters": f,
                       "prompt": "Explain this view: Region = South, Jun 2024"}),
        _click("pdf", {"sheet": "Monthly Sales", "style": "explore", "measure": "Revenue", "dim": "Region", "filters": json.dumps(f)}),
    ])
    msgs = _messages(turns[1])
    validate.check(msgs)
    assert "Explain this view: Region = South" in fake.calls[-1][0] and "filtered to Region = South" in fake.calls[-1][0]
    comps = _components(msgs)
    assert _explorer_data(comps["frame"]["htmlContent"])["init"]
    assert json.loads(comps["pdf"]["action"]["event"]["context"]["filters"]) == f  # buttons carry filters as JSON
    files = [p for e in turns[2] for p in e.content.parts if p.inline_data and p.inline_data.mime_type == "application/pdf"]
    assert len(files) == 1 and files[0].inline_data.data.startswith(b"%PDF")


@pytest.mark.parametrize("style", ["explore", "heatmap", "trend"])
def test_every_view_has_a_theme_switch_and_themed_heatmap_colours(tabs, style):
    a = analysis.analyze(_tab(tabs, "Monthly Sales"), style)
    html_ = page.page(a, llm.fallback(a), BOOK.name)
    assert all(f'data-t="{t}"' in html_ for t in ("auto", "light", "dark"))
    assert ":root[data-theme=dark]" in html_ and ":root:not([data-theme=light])" in html_
    assert not re.search(r'style="background:#', html_)  # cell colours come from the theme, not fixed hex


def test_one_row_per_category_sheet_charts_in_every_style():
    # A small weather sheet: every city is its own row, so there are no peer groups for anomalies.
    df = pd.DataFrame({"city": ["London", "Paris", "Delhi", "Mumbai", "Tokyo", "Sydney", "Cairo", "Oslo"],
                       "temperature highest": ["18°C", "21°C", "38°C", "33°C", "26°C", "24°C", "35°C", "9°C"],
                       "temperature lowest": [9, 12, 27, 26, 18, 14, 22, 1],
                       "humidity highest": ["88%", "80%", "60%", "90%", "75%", "70%", "40%", "85%"],
                       "humidity lowest": [55, 50, 25, 70, 45, 40, 15, 60]})
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Weather")
    up = sheets.Upload("weather.xlsx", buf.getvalue(), sheets.XLSX)
    sheet, = sheets.read(up)
    assert set(sheets.profile(sheet).measures) == {"temperature highest", "temperature lowest", "humidity highest", "humidity lowest"}
    for style in analysis.STYLES:
        a = analysis.analyze(sheet, style)
        validate.check(render.dashboard_card(a, llm.fallback(a), up.name))


def _sheet_of(df: pd.DataFrame) -> sheets.Sheet:
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    sheet, = sheets.read(sheets.Upload("x.xlsx", buf.getvalue(), sheets.XLSX))
    return sheet


_rng = np.random.default_rng(1)
_CITIES = ["London", "Paris", "Delhi", "Mumbai", "Tokyo", "Sydney", "Cairo", "Oslo", "Lima", "Rome", "Seoul", "Dubai"]
SHAPES = {
    "numbers only": pd.DataFrame({"a": _rng.normal(size=30), "b": _rng.normal(size=30)}),
    "one number column": pd.DataFrame({"v": _rng.integers(1, 100, 40)}),
    "dates and one number": pd.DataFrame({"date": pd.to_datetime(["2024-01-01", "2024-02-01"]), "v": [1, 2]}),
    "year column": pd.DataFrame({"year": [2020, 2021, 2022, 2023], "v": [1, 4, 9, 16]}),
    "one number per city": pd.DataFrame({"city": _CITIES, "v": _rng.integers(1, 50, 12)}),
    "two rows": pd.DataFrame({"city": ["A", "B"], "v": [1, 2]}),
    "gaps and text": pd.DataFrame({"city": _CITIES, "v": [1, None, 3, "n/a", 5, 6, "-", 8, 9, None, 11, 12], "w": [None] * 12}),
    "single period": pd.DataFrame({"date": pd.to_datetime(["2024-01-01"] * 10), "city": _CITIES[:10], "v": range(10)}),
}


@pytest.mark.parametrize("shape", SHAPES)
def test_unusual_sheet_shapes_chart_in_every_style(shape):
    sheet = _sheet_of(SHAPES[shape])
    for style in analysis.STYLES:
        a = analysis.analyze(sheet, style)
        messages = render.dashboard_card(a, llm.fallback(a), "x.xlsx")
        validate.check(messages)  # e.g. no null values in button contexts when there is no category column
        if a.style == "explore":
            assert _explorer_data(page.page(a, llm.fallback(a), "x.xlsx"))["cube"]["dims"]


@pytest.mark.parametrize("df,message", [
    (pd.DataFrame({"city": ["London"], "v": [18]}), "only one row"),
    (pd.DataFrame({"city": _CITIES, "v": [0] * 12}), "same value in every row"),
])
def test_sheets_with_nothing_to_compare_say_why(df, message):
    with pytest.raises(ValueError, match=message):
        analysis.analyze(_sheet_of(df))


def test_title_above_the_table_does_not_leave_an_empty_column():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Weather report Q3"
    ws.append([])
    ws.append([None, "City", "Max temp", "Min temp"])
    for row in [("London", 18, 9), ("Paris", 21, 12), ("Rome", 29, 19), ("Oslo", 9, 1)]:
        ws.append([None, *row])
    buf = io.BytesIO()
    wb.save(buf)
    sheet, = sheets.read(sheets.Upload("x.xlsx", buf.getvalue(), sheets.XLSX))
    assert list(sheet.df.columns) == ["City", "Max temp", "Min temp"]


# --- upfront file review ------------------------------------------------------------------

def _messy_book() -> bytes:
    import datetime as dt
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"], ws["A2"] = "Q3 sales report", "Prepared by finance"
    ws.append([])
    ws.append(["Date", "Region", "Product", "Revenue", "Units", None, "Units"])
    regions = ["North", "South", "East", "north ", "West"]
    for i in range(40):
        rev = {5: "n/a", 9: "#DIV/0!", 12: None}.get(i, 1000 + i * 37)
        ws.append([dt.date(2024, 1 + i % 9, 1 + i % 27), regions[i % 5], "Alpha" if i % 2 else "Beta", rev, i % 7 + 1,
                   3 if i % 4 == 0 else None, i])
    ws.append(["2024-09-01", "South", "Beta", 1500, 2, None, 2])
    ws.append(["2024-09-01", "South", "Beta", 1500, 2, None, 2])
    ws.append(["Total", None, None, 99999, None, None, None])
    ws.append([])
    ws.append(["Target", "Owner"])
    ws.append(["North", "Ann"])
    wb.create_sheet("Notes")["A1"] = "Remember to update monthly"
    wb.create_sheet("Blank")
    ws3 = wb.create_sheet("Customers")
    ws3.append(["Customer", "Spend"])
    for i in range(10):
        ws3.append([f"C{i}", 100 + i if i < 6 else f"x{i}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MESSY = sheets.Upload("sales_q3.xlsx", _messy_book(), sheets.XLSX)


def test_review_points_at_cells_and_says_how_to_fix():
    from agent import review
    tabs = {t.name: t for t in review.review(MESSY)}
    sales = [f.title for f in tabs["Sales"].findings]
    want = ["Column F has values but no header", "2 columns are called “Units” (E, G)", "second table below the blank row 48",
            "Revenue: 1 formula error (D14 “#DIV/0!”)", "Revenue: 1 cell isn't a number (D10 “n/a”)",
            "(“North” / “north”)", "exact copy of an earlier row (row 46)", "table starts at row 4", "Left out 1 total row"]
    for w in want:
        assert any(w in t for t in sales), w
    assert all(f.tip for f in tabs["Sales"].findings if f.level != "info")
    assert tabs["Sales"].status == "fix" and tabs["Sales"].chartable
    assert tabs["Blank"].status == "empty" and tabs["Notes"].status == "blocked"
    assert tabs["Customers"].status == "blocked" and "read as text" in tabs["Customers"].findings[0].title
    assert review.counts(list(tabs.values())) == (5, 4)


def test_review_of_the_clean_sample_finds_nothing_to_fix():
    from agent import review
    tabs = review.review(BOOK)
    assert review.counts(tabs)[0] == 0 and all(t.chartable for t in tabs)
    assert "looks clean" in review.headline(BOOK.name, tabs) or "worth a look" in review.headline(BOOK.name, tabs)


def test_review_of_a_csv_uses_line_numbers():
    from agent import review
    up = sheets.Upload("w.csv", b"city,temp\nLondon,18\nParis,n/a\nRome,29\nOslo,9\nLima,20\n", sheets.CSV)
    t, = review.review(up)
    assert any("line 3 “n/a”" in f.title for f in t.findings)


def test_upload_offers_review_with_a_hint_and_review_opens_the_check(fake):
    turns = _converse([_upload(MESSY), _click("review", {"prompt": "Review my file"}), _text("can you check my file?")])
    first = _components(_messages(turns[0]))
    assert first["review"]["action"]["event"]["name"] == "review"
    assert "5 things to fix" in first["reviewd"]["text"]
    for t in turns[1:]:
        msgs = _messages(t)
        validate.check(msgs)
        comps = _components(msgs)
        assert comps["root"]["component"] == "Canvas" and "Column F has values but no header" in comps["frame"]["htmlContent"]
    assert not fake.calls  # the review needs no model call


def test_file_with_no_numbers_offers_the_review():
    up = sheets.Upload("names.csv", b"name,team\nAnn,Red\nBo,Blue\nCy,Red\n", sheets.CSV)
    turns = _converse([_upload(up)])
    comps = _components(_messages(turns[0]))
    assert comps["review"]["action"]["event"]["name"] == "review"


def test_welcome_card_validates():
    validate.check(render.welcome_card())
